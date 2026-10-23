from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - documented setup guard
    raise SystemExit(
        "Falta openpyxl. Ejecute este script con el Python incluido en Codex o instale openpyxl."
    ) from exc


PROJECT = Path(__file__).resolve().parents[1]
WORKSPACE = PROJECT.parent
SOURCE = WORKSPACE / "Ghagpt"
RAW = SOURCE / "data" / "raw"
PUBLIC_DATA = PROJECT / "public" / "data"
LIBRARY = PROJECT / "fuente_de_datos"

FOLDER_MAP = {
    "amsclae": "AMSCLAE",
    "fao_aquastat": "FAO_AQUASTAT",
    "icc": "ICC",
    "ine_2022": "INE",
    "ine_2024": "INE",
    "ine_amsa": "AMSA",
    "ine_amsclae": "AMSCLAE",
    "ine_insivumeh": "INSIVUMEH",
    "ine_mspas": "MSPAS",
    "nasa_power": "NASA_POWER",
    "unep_gemstat": "GEMStat",
    "usac": "USAC",
    "water_quality_portal": "Water_Quality_Portal",
    "world_bank": "Banco_Mundial",
    "zenodo_uvg": "UVG_Zenodo",
}

REQUIRED_FOLDERS = [
    "INE",
    "INSIVUMEH",
    "MSPAS",
    "AMSA",
    "AMSCLAE",
    "UVG_Zenodo",
    "USAC",
    "ICC",
    "NASA_POWER",
    "Banco_Mundial",
    "GEMStat",
    "FAO_AQUASTAT",
    "Water_Quality_Portal",
    "otras_fuentes",
]

SELECTED = {
    "GTW-004": "Cobertura de calidad en Amatitlán y la cuenca del Villalobos, con amplia diversidad de parámetros.",
    "GTW-006": "Monitoreo de ríos tributarios de Atitlán con caudal, fisicoquímica, nutrientes e indicadores microbiológicos.",
    "GTW-025": "Evidencia académica sobre contaminantes emergentes en Las Vacas y su conexión con el Motagua.",
    "GTW-028": "La serie continua más extensa del paquete para un cuerpo de agua guatemalteco: Atitlán, 2010-2024.",
    "GTW-031": "Serie meteorológica diaria de apoyo para Las Vacas, 2001-2025, sin requerir una clave privada.",
    "GTW-033": "Serie meteorológica diaria en la desembocadura del Motagua para contrastar la cuenca alta y baja.",
    "GTW-045": "Lectura territorial de los ríos de la costa sur producida por ICC, con continuidad de informes 2017-2026.",
    "GTW-052": "Estación georreferenciada real en el lago de Atitlán y resultados públicos del Water Quality Portal.",
    "GTW-067": "Producción mensual de agua potable de EMPAGUA en 2024, útil para la dimensión de abastecimiento urbano.",
    "GTW-068": "Monitoreo oficial 2024 de INSIVUMEH con coordenadas y parámetros fisicoquímicos para el mapa inicial.",
    "GTW-069": "Calidad del agua de AMSA en 2024, representativa del lago de Amatitlán y la cuenca del Villalobos.",
    "GTW-072": "Vigilancia municipal de cloro residual de MSPAS en 2024, con cobertura nacional agregada.",
}


def read_csv(path: Path, delimiter: str = ",") -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def slug_topic(value: str) -> str:
    text = value.lower()
    if any(key in text for key in ("calidad", "monitoreo", "microbi", "cloro", "limnolog")):
        return "Calidad del agua"
    if any(key in text for key in ("caudal", "hidrolog", "río", "rio")):
        return "Hidrología"
    if any(key in text for key in ("potable", "abastecimiento", "comunidades", "sistema")):
        return "Abastecimiento"
    if any(key in text for key in ("tratamiento", "ptar", "saneamiento", "residual")):
        return "Saneamiento"
    if any(key in text for key in ("meteor", "precip", "clima", "power")):
        return "Clima"
    if any(key in text for key in ("subterránea", "subterranea", "pozo", "riego")):
        return "Agua subterránea"
    if any(key in text for key in ("plástico", "plastico", "contamin", "fármaco", "farmaco")):
        return "Contaminación"
    return "Referencia técnica"


def access_state(master: dict[str, str]) -> str:
    access = (master.get("estado_acceso") or "").lower()
    downloaded = (master.get("estado_descarga") or "").lower()
    if downloaded == "descargado":
        return "Descargado y verificado"
    if "solicitud" in access or "registro" in access:
        return "Requiere solicitud o registro"
    if "guatemala" in access and "cero" in access:
        return "Sin cobertura para Guatemala"
    if "públic" in access or "localizado" in access:
        return "Localizado"
    return "Pendiente"


def institution_folder(local_path: str) -> str:
    parts = Path(local_path.replace("/", "\\")).parts
    source_folder = parts[2] if len(parts) > 2 else ""
    return FOLDER_MAP.get(source_folder, "otras_fuentes")


def human_size(byte_count: int) -> str:
    units = ["B", "KB", "MB", "GB"]
    value = float(byte_count)
    for unit in units:
        if value < 1024 or unit == units[-1]:
            return f"{value:.0f} {unit}" if unit == "B" else f"{value:.1f} {unit}"
        value /= 1024
    return f"{byte_count} B"


def parse_number(value: Any) -> float | None:
    if value in (None, "", "-", "NA", "N/A"):
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def safe_params(value: str, limit: int = 8) -> list[str]:
    parts = [re.sub(r"\s+", " ", item).strip() for item in (value or "").split("|")]
    return [item for item in parts if item][:limit]


def build_library() -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    masters = read_csv(SOURCE / "inventories" / "inventario_maestro.csv")
    profiles = read_csv(SOURCE / "inventories" / "archivos_descargados.csv")
    by_path = {row["ruta_local"].replace("\\", "/"): row for row in masters if row.get("ruta_local")}

    for folder in REQUIRED_FOLDERS:
        (LIBRARY / folder).mkdir(parents=True, exist_ok=True)

    catalog: list[dict[str, Any]] = []
    manifest: list[dict[str, str]] = []

    for index, profile in enumerate(profiles, start=1):
        local_path = profile["local_path"].replace("\\", "/")
        source_path = SOURCE / Path(local_path)
        if not source_path.exists():
            raise FileNotFoundError(f"No existe el archivo inventariado: {source_path}")
        master = by_path.get(local_path)
        if not master:
            raise KeyError(f"No existe ficha maestra para {local_path}")

        folder = institution_folder(local_path)
        destination = LIBRARY / folder / source_path.name
        shutil.copy2(source_path, destination)
        relative_copy = destination.relative_to(PROJECT).as_posix()
        source_url = master.get("url_fuente") or master.get("url_descarga") or ""
        byte_count = int(profile.get("bytes") or source_path.stat().st_size)

        manifest.append(
            {
                "numero": str(index),
                "archivo_original": local_path,
                "copia": relative_copy,
                "institucion": master.get("organizacion", ""),
                "url_procedencia": source_url,
                "sha256": profile.get("sha256", ""),
            }
        )

        title = master.get("titulo") or profile.get("file_name") or source_path.name
        catalog.append(
            {
                "number": index,
                "id": master.get("id_recurso") or f"FILE-{index:03d}",
                "title": title,
                "fileName": source_path.name,
                "institution": master.get("organizacion", "Sin institución"),
                "topic": slug_topic(" ".join((title, master.get("categoria", ""), master.get("parametros_principales", "")))),
                "territory": master.get("cobertura_geografica") or "No especificado",
                "period": master.get("cobertura_temporal") or profile.get("date_min") or "No especificado",
                "format": (master.get("formato") or profile.get("extension") or "").upper(),
                "bytes": byte_count,
                "size": human_size(byte_count),
                "accessState": access_state(master),
                "summary": master.get("parametros_principales") or master.get("notas") or "Consulte la ficha técnica del recurso.",
                "sourceUrl": source_url,
                "downloadUrl": master.get("url_descarga") or "",
                "localPath": f"../{relative_copy}",
                "license": master.get("licencia") or "No indicada",
                "sha256": profile.get("sha256") or "",
                "records": profile.get("records") or master.get("registros_o_paginas") or "",
                "dateMin": profile.get("date_min") or "",
                "dateMax": profile.get("date_max") or "",
                "parameters": safe_params(profile.get("parameters") or master.get("parametros_principales", "")),
                "coordinates": profile.get("coordinates_present") or master.get("coordenadas") or "unknown",
                "selectedMvp": (master.get("id_recurso") or "") in SELECTED,
            }
        )

    manifest_path = LIBRARY / "manifest_fuentes.csv"
    with manifest_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(manifest[0].keys()))
        writer.writeheader()
        writer.writerows(manifest)

    if len(catalog) != 75:
        raise ValueError(f"Se esperaban 75 archivos verificados y se procesaron {len(catalog)}")
    return catalog, manifest


def feature(
    feature_id: str,
    lon: float,
    lat: float,
    *,
    name: str,
    kind: str,
    institution: str,
    source_ids: list[str],
    latest: str,
    basin: str = "",
    department: str = "",
    municipality: str = "",
    parameters: Iterable[str] = (),
    evidence: Iterable[dict[str, Any]] = (),
    series: Iterable[dict[str, Any]] = (),
    confidence: str = "Alta",
    warning: str = "La clasificación técnica requiere más series y control de calidad.",
) -> dict[str, Any]:
    return {
        "type": "Feature",
        "id": feature_id,
        "geometry": {"type": "Point", "coordinates": [round(lon, 6), round(lat, 6)]},
        "properties": {
            "name": name,
            "kind": kind,
            "basin": basin,
            "department": department,
            "municipality": municipality,
            "emotion": "Sin datos suficientes",
            "emotionCode": "insufficient",
            "emotionReason": "El MVP no convierte una campaña o serie parcial en una conclusión científica.",
            "latest": latest,
            "parameters": list(parameters),
            "institution": institution,
            "sourceIds": source_ids,
            "confidence": confidence,
            "warning": warning,
            "evidence": list(evidence),
            "series": list(series),
        },
    }


def load_insivumeh_2023() -> list[dict[str, Any]]:
    rows = read_csv(RAW / "ine_insivumeh" / "calidad_agua_insivumeh_2023.csv", ";")
    groups: dict[tuple[str, str, str], list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        if parse_number(row.get("latitud")) is None or parse_number(row.get("longitud")) is None:
            continue
        groups[(row.get("punto_monitoreo", ""), row.get("latitud", ""), row.get("longitud", ""))].append(row)

    output = []
    for index, ((name, lat_raw, lon_raw), items) in enumerate(groups.items(), start=1):
        params = list(dict.fromkeys(re.sub(r"\s+", " ", row.get("atributo", "")).strip() for row in items if row.get("atributo")))
        evidence = []
        for row in items:
            value = parse_number(row.get("valor"))
            if value is not None and len(evidence) < 5:
                evidence.append(
                    {
                        "parameter": re.sub(r"\s+", " ", row.get("atributo", "")).strip(),
                        "value": value,
                        "unit": "Unidad incluida en el nombre del parámetro",
                        "date": row.get("fecha_muestreo") or "2023",
                        "station": name,
                        "sourceId": "GTW-003",
                    }
                )
        first = items[0]
        output.append(
            feature(
                f"insivumeh-2023-{index}",
                float(lon_raw),
                float(lat_raw),
                name=name or first.get("rio") or "Punto INSIVUMEH",
                kind="Punto de muestreo",
                institution="INE / INSIVUMEH",
                source_ids=["GTW-003"],
                latest=max((row.get("fecha_muestreo", "") for row in items), default="2023"),
                basin=first.get("cuenca", ""),
                department=first.get("depto", ""),
                municipality=first.get("municipio", ""),
                parameters=params[:10],
                evidence=evidence,
                confidence="Alta para ubicación; media para cobertura temporal",
                warning="Datos de 2023. No representan condiciones actuales ni una evaluación integral del cuerpo de agua.",
            )
        )
    return output


def load_insivumeh_2024() -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(
        RAW / "ine_2024" / "4._Base_de_monitoreo_del_Agua_INSIVUMEH_2024.xlsx",
        read_only=True,
        data_only=True,
    )
    worksheet = workbook["Base"]
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    records = [dict(zip(headers, values)) for values in rows]
    groups: dict[tuple[str, float, float], list[dict[str, Any]]] = defaultdict(list)
    for row in records:
        lat = parse_number(row.get("latitud"))
        lon = parse_number(row.get("longitud"))
        if lat is not None and lon is not None:
            groups[(str(row.get("punto_monitoreo") or ""), lat, lon)].append(row)

    output = []
    for index, ((name, lat, lon), items) in enumerate(groups.items(), start=1):
        params = list(dict.fromkeys(re.sub(r"\s+", " ", str(row.get("parametro") or "")).strip() for row in items))
        evidence = []
        for row in items:
            value = parse_number(row.get("valor"))
            if value is not None and len(evidence) < 6:
                evidence.append(
                    {
                        "parameter": re.sub(r"\s+", " ", str(row.get("parametro") or "")).strip(),
                        "value": value,
                        "unit": str(row.get("unidad_medida") or "No indicada"),
                        "date": f"{row.get('mes') or ''} 2024".strip(),
                        "station": name,
                        "sourceId": "GTW-068",
                    }
                )
        first = items[0]
        output.append(
            feature(
                f"insivumeh-2024-{index}",
                lon,
                lat,
                name=name or str(first.get("cuerpo_agua") or "Punto INSIVUMEH"),
                kind="Punto de monitoreo",
                institution="INE / INSIVUMEH",
                source_ids=["GTW-068"],
                latest="2024",
                basin=str(first.get("cuenca") or ""),
                department=str(first.get("departamento") or ""),
                municipality=str(first.get("municipio") or ""),
                parameters=params[:10],
                evidence=evidence,
                confidence="Alta para ubicación y procedencia; media para cobertura anual",
                warning="La base disponible para 2024 no constituye una serie histórica continua por punto.",
            )
        )
    return output


def read_nasa(path: Path, source_id: str, name: str, basin: str) -> tuple[dict[str, Any], dict[str, Any]]:
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    location = next(line for line in lines if line.startswith("Location:"))
    match = re.search(r"latitude\s+(-?\d+(?:\.\d+)?)\s+longitude\s+(-?\d+(?:\.\d+)?)", location)
    if not match:
        raise ValueError(f"No se pudo leer la ubicación NASA POWER de {path}")
    lat, lon = float(match.group(1)), float(match.group(2))
    header_end = lines.index("-END HEADER-")
    reader = csv.DictReader(lines[header_end + 1 :])
    by_year: dict[int, list[float]] = defaultdict(list)
    row_count = 0
    for row in reader:
        row_count += 1
        precipitation = parse_number(row.get("PRECTOTCORR"))
        if precipitation is not None and precipitation > -900:
            by_year[int(row["YEAR"])].append(precipitation)
    series = [
        {"label": str(year), "value": round(sum(values), 1), "unit": "mm/año"}
        for year, values in sorted(by_year.items())[-10:]
    ]
    map_feature = feature(
        source_id.lower(),
        lon,
        lat,
        name=name,
        kind="Estación meteorológica de apoyo",
        institution="NASA POWER",
        source_ids=[source_id],
        latest="2025-12-31",
        basin=basin,
        parameters=["Precipitación", "Temperatura a 2 m", "Humedad relativa", "Viento a 2 m"],
        evidence=[
            {
                "parameter": "Precipitación anual",
                "value": series[-1]["value"],
                "unit": "mm/año",
                "date": series[-1]["label"],
                "station": name,
                "sourceId": source_id,
            }
        ],
        series=series,
        confidence="Alta para la coordenada consultada; es una celda modelada, no una estación en terreno",
        warning="NASA POWER aporta variables meteorológicas modeladas. No mide calidad de agua ni caudal del río.",
    )
    stats = {"records": row_count, "start": min(by_year), "end": max(by_year), "series": series}
    return map_feature, stats


def load_wqp() -> list[dict[str, Any]]:
    station_rows = read_csv(RAW / "water_quality_portal" / "atitlan_nalms_station.csv")
    result_rows = read_csv(RAW / "water_quality_portal" / "atitlan_nalms_results.csv")
    output = []
    for index, station in enumerate(station_rows, start=1):
        station_id = station.get("MonitoringLocationIdentifier", "")
        matching = [row for row in result_rows if row.get("MonitoringLocationIdentifier") == station_id]
        parameters = list(dict.fromkeys(row.get("CharacteristicName", "") for row in matching if row.get("CharacteristicName")))
        dates = [row.get("ActivityStartDate", "") for row in matching if row.get("ActivityStartDate")]
        evidence = []
        for row in matching:
            value = parse_number(row.get("ResultMeasureValue"))
            if value is not None and len(evidence) < 5:
                evidence.append(
                    {
                        "parameter": row.get("CharacteristicName") or "Resultado",
                        "value": value,
                        "unit": row.get("ResultMeasure/MeasureUnitCode") or "No indicada",
                        "date": row.get("ActivityStartDate") or "Sin fecha",
                        "station": station.get("MonitoringLocationName") or station_id,
                        "sourceId": "GTW-052",
                    }
                )
        output.append(
            feature(
                f"wqp-atitlan-{index}",
                float(station["LongitudeMeasure"]),
                float(station["LatitudeMeasure"]),
                name=station.get("MonitoringLocationName") or "Lago de Atitlán",
                kind="Lago / punto de observación",
                institution=station.get("OrganizationFormalName") or "Water Quality Portal",
                source_ids=["GTW-051", "GTW-052"],
                latest=max(dates) if dates else "Sin fecha",
                basin="Lago de Atitlán",
                department="Sololá",
                parameters=parameters,
                evidence=evidence,
                confidence="Media: coordenada publicada con método y datum marcados como desconocidos por el proveedor",
                warning="La ficha del proveedor no declara método ni datum horizontal; se conserva la coordenada sin reinterpretarla.",
            )
        )
    return output


def summarize_uvg() -> dict[str, Any]:
    rows = read_csv(RAW / "zenodo_uvg" / "Lake_Atitlan_water_quality_CWG_2010-2024.csv")
    yearly: dict[str, list[float]] = defaultdict(list)
    depths: list[float] = []
    for row in rows:
        year = row.get("date", "")[:4]
        oxygen = parse_number(row.get("dissolved_oxygen"))
        depth = parse_number(row.get("depth"))
        if year and oxygen is not None:
            yearly[year].append(oxygen)
        if depth is not None:
            depths.append(depth)
    series = [
        {"label": year, "value": round(sum(values) / len(values), 2), "unit": "mg/L"}
        for year, values in sorted(yearly.items())
    ]
    return {
        "records": len(rows),
        "dateMin": min(row["date"] for row in rows if row.get("date")),
        "dateMax": max(row["date"] for row in rows if row.get("date")),
        "maxDepth": max(depths) if depths else None,
        "series": series,
        "seriesLabel": "Oxígeno disuelto medio por observación y año (descriptivo; no ajustado por profundidad)",
    }


def summarize_mspas_2024() -> dict[str, Any]:
    workbook = openpyxl.load_workbook(
        RAW / "ine_2024" / "8._Base_de_Vigilancia_cloro_residual_MSPAS_2024.xlsx",
        read_only=True,
        data_only=True,
    )
    rows = workbook[workbook.sheetnames[0]].iter_rows(values_only=True)
    headers = [str(value or "") for value in next(rows)]
    records = [dict(zip(headers, values)) for values in rows]
    total = sum(parse_number(row.get("total_de_muestras")) or 0 for row in records)
    inside = sum(parse_number(row.get("muestras_dentro_de_rango")) or 0 for row in records)
    outside = sum(parse_number(row.get("muestras_fuera_de_rango")) or 0 for row in records)
    return {
        "records": len(records),
        "totalSamples": int(total),
        "insideRange": int(inside),
        "outsideRange": int(outside),
        "insideShare": round(inside / total * 100, 1) if total else None,
    }


def build_selected(catalog: list[dict[str, Any]], nasa_stats: dict[str, Any]) -> list[dict[str, Any]]:
    uvg = summarize_uvg()
    mspas = summarize_mspas_2024()
    by_id = {item["id"]: item for item in catalog}
    details: dict[str, dict[str, Any]] = {
        "GTW-028": uvg,
        "GTW-072": mspas,
        **nasa_stats,
    }
    selected = []
    for source_id, reason in SELECTED.items():
        item = dict(by_id[source_id])
        item["selectionReason"] = reason
        item["profile"] = details.get(source_id, {})
        selected.append(item)
    return selected


def build_catalog_html(catalog: list[dict[str, Any]]) -> None:
    rows = []
    for item in catalog:
        source_link = (
            f'<a href="{item["sourceUrl"]}" target="_blank" rel="noreferrer">Procedencia</a>'
            if item["sourceUrl"]
            else "Sin URL"
        )
        rows.append(
            "<tr>"
            f'<td data-label="#">{item["number"]}</td>'
            f'<td data-label="Recurso"><strong>{html_escape(item["title"])}</strong><small>{html_escape(item["fileName"])}</small></td>'
            f'<td data-label="Institución">{html_escape(item["institution"])}</td>'
            f'<td data-label="Tema">{html_escape(item["topic"])}</td>'
            f'<td data-label="Territorio">{html_escape(item["territory"])}</td>'
            f'<td data-label="Periodo">{html_escape(item["period"])}</td>'
            f'<td data-label="Formato">{html_escape(item["format"])} · {html_escape(item["size"])}</td>'
            f'<td data-label="Acceso">{html_escape(item["accessState"])}</td>'
            f'<td data-label="Resumen">{html_escape(item["summary"])}</td>'
            f'<td data-label="Enlaces"><a href="{item["localPath"].removeprefix("../")}">Archivo local</a>{source_link}</td>'
            f'<td data-label="Licencia">{html_escape(item["license"])}</td>'
            "</tr>"
        )
    html = f"""<!doctype html>
<html lang="es">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Catálogo maestro de fuentes de agua de Guatemala</title>
  <style>
    @font-face {{ font-family: Atlas; src: url('public/fonts/UbuntuCondensed.ttf'); }}
    @font-face {{ font-family: AtlasText; src: url('public/fonts/Ubuntu-Regular.ttf'); }}
    :root {{ color-scheme: light; --ink:#0e2d36; --water:#0a706c; --paper:#edf5f3; --rule:#91adae; --alert:#b34238; }}
    * {{ box-sizing:border-box; }} body {{ margin:0; background:var(--paper); color:var(--ink); font:15px/1.45 AtlasText,Arial,sans-serif; }}
    header {{ padding:34px clamp(18px,4vw,64px) 28px; color:#fbfdfc; background:var(--ink); border-bottom:6px solid var(--water); }}
    h1 {{ margin:0; max-width:900px; font:400 clamp(38px,6vw,74px)/.92 Atlas,Arial Narrow,sans-serif; }}
    header p {{ max-width:78ch; color:#d7e4e4; }} .meta {{ display:flex; gap:24px; flex-wrap:wrap; font-family:Atlas; text-transform:uppercase; letter-spacing:.08em; }}
    main {{ padding:24px clamp(12px,3vw,44px) 60px; }} table {{ width:100%; border-collapse:collapse; background:#fbfdfc; }}
    th {{ position:sticky; top:0; z-index:1; background:var(--water); color:white; font:400 14px Atlas; letter-spacing:.04em; text-align:left; }}
    th,td {{ padding:10px; border:1px solid var(--rule); vertical-align:top; }} tbody tr:nth-child(even) {{ background:#dfecea; }}
    small {{ display:block; margin-top:4px; color:#45656b; word-break:break-word; }} a {{ display:block; color:var(--water); font-weight:700; text-underline-offset:3px; }}
    a:focus-visible {{ outline:3px solid var(--alert); outline-offset:2px; }}
    @media(max-width:900px) {{ thead {{ display:none; }} tr {{ display:block; margin-bottom:18px; border:1px solid var(--ink); }} td {{ display:grid; grid-template-columns:120px 1fr; border-width:0 0 1px; }} td::before {{ content:attr(data-label); font-family:Atlas; text-transform:uppercase; letter-spacing:.08em; }} }}
    @media print {{ header {{ background:white; color:var(--ink); }} header p {{ color:var(--ink); }} th {{ position:static; }} main {{ padding:0; }} body {{ font-size:9px; }} }}
  </style>
</head>
<body>
  <header>
    <h1>Catálogo maestro de fuentes de agua de Guatemala</h1>
    <p>Índice navegable de los 75 archivos descargados y verificados. Los enlaces locales son relativos a la carpeta del proyecto; algunos navegadores pueden requerir abrir este archivo desde el disco.</p>
    <div class="meta"><span>75 archivos</span><span>Fecha de corte: 21 oct 2026</span><span>Integridad: SHA-256 en fichas técnicas</span></div>
  </header>
  <main><table><thead><tr><th>#</th><th>Recurso</th><th>Institución</th><th>Tema</th><th>Territorio</th><th>Periodo</th><th>Formato</th><th>Acceso</th><th>Resumen</th><th>Enlaces</th><th>Licencia</th></tr></thead><tbody>{''.join(rows)}</tbody></table></main>
</body></html>"""
    (PROJECT / "catalogo_fuentes_agua_guatemala.html").write_text(html, encoding="utf-8")


def html_escape(value: Any) -> str:
    return (
        str(value or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def verify_copies(catalog: list[dict[str, Any]]) -> dict[str, Any]:
    failures = []
    for item in catalog:
        target = PROJECT / item["localPath"].removeprefix("../")
        if not target.exists():
            failures.append({"id": item["id"], "error": "missing", "path": str(target)})
            continue
        digest = hashlib.sha256(target.read_bytes()).hexdigest()
        if digest != item["sha256"]:
            failures.append({"id": item["id"], "error": "sha256", "path": str(target)})
    return {"checked": len(catalog), "failures": failures, "passed": not failures}


def main() -> None:
    PUBLIC_DATA.mkdir(parents=True, exist_ok=True)
    fonts = PROJECT / "public" / "fonts"
    fonts.mkdir(parents=True, exist_ok=True)
    for font in ("UbuntuCondensed.ttf", "Ubuntu-Regular.ttf", "Ubuntu-Bold.ttf"):
        shutil.copy2(SOURCE / "report" / "assets" / font, fonts / font)

    catalog, manifest = build_library()
    features = load_insivumeh_2023() + load_insivumeh_2024() + load_wqp()
    nasa_stats: dict[str, Any] = {}
    nasa_specs = [
        ("nasa_power_las_vacas_daily_2001_2025.csv", "GTW-031", "Las Vacas · celda NASA POWER", "Río Las Vacas / Motagua"),
        ("nasa_power_motagua_middle_daily_2001_2025.csv", "GTW-032", "Motagua medio · celda NASA POWER", "Río Motagua"),
        ("nasa_power_motagua_outlet_daily_2001_2025.csv", "GTW-033", "Desembocadura del Motagua · celda NASA POWER", "Río Motagua"),
    ]
    for filename, source_id, name, basin in nasa_specs:
        map_feature, stats = read_nasa(RAW / "nasa_power" / filename, source_id, name, basin)
        features.append(map_feature)
        nasa_stats[source_id] = stats

    selected = build_selected(catalog, nasa_stats)
    formats = sorted({item["format"] for item in catalog})
    institutions = sorted({item["institution"] for item in catalog})
    topics = sorted({item["topic"] for item in catalog})
    access_states = sorted({item["accessState"] for item in catalog})
    summary = {
        "cutoff": "2026-10-21",
        "inventoriedResources": 113,
        "downloadedFiles": 75,
        "organizations": 42,
        "missingDataRecords": 16,
        "selectedMvpDatasets": len(selected),
        "mappedFeatures": len(features),
        "formats": formats,
        "institutions": institutions,
        "topics": topics,
        "accessStates": access_states,
        "mapStyle": "https://tiles.openfreemap.org/styles/dark",
        "mapAttribution": "OpenFreeMap © OpenMapTiles · Datos de OpenStreetMap",
    }

    write_json(PUBLIC_DATA / "catalog.json", catalog)
    write_json(PUBLIC_DATA / "selected_datasets.json", selected)
    write_json(PUBLIC_DATA / "water_features.geojson", {"type": "FeatureCollection", "features": features})
    write_json(PUBLIC_DATA / "summary.json", summary)
    write_json(PROJECT / "verification" / "copy_integrity.json", verify_copies(catalog))
    build_catalog_html(catalog)
    public_docs = PROJECT / "public" / "docs"
    public_docs.mkdir(parents=True, exist_ok=True)
    for document in ("DATA_METHOD.md", "EMOTION_MODEL.md", "MVP_DATASETS.md"):
        document_path = PROJECT / document
        if document_path.exists():
            shutil.copy2(document_path, public_docs / document)
    shutil.copy2(PROJECT / "catalogo_fuentes_agua_guatemala.html", PROJECT / "public" / "catalogo_fuentes_agua_guatemala.html")
    print(f"Catálogo: {len(catalog)} archivos")
    print(f"Selección MVP: {len(selected)} conjuntos")
    print(f"Elementos cartográficos: {len(features)}")
    print(f"Manifiesto: {len(manifest)} filas")


if __name__ == "__main__":
    main()
